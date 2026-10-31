"""Build ANLEITUNG_DE.docx from docs/ANLEITUNG_DE.md (subset of Markdown).

    python examples/build_anleitung_docx.py          # DOCX only
    python examples/build_anleitung_docx.py --odt    # DOCX + ODT (LibreOffice)

The ODT step needs a LibreOffice install; without it the DOCX is still written
and the conversion is skipped with a message.
"""

from __future__ import annotations

import argparse
import re
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
MD_PATH = ROOT / "docs" / "ANLEITUNG_DE.md"
DOCX_PATH = ROOT / "docs" / "ANLEITUNG_DE.docx"

SOFFICE_CANDIDATES = (
    r"C:\Program Files\LibreOffice\program\soffice.exe",
    r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
    "/usr/bin/soffice",
    "/usr/local/bin/soffice",
)


def _find_soffice() -> str | None:
    exe = shutil.which("soffice") or shutil.which("libreoffice")
    if exe:
        return exe
    for cand in SOFFICE_CANDIDATES:
        if Path(cand).exists():
            return cand
    return None


def convert_to_odt(docx_path: Path) -> Path | None:
    """Convert ``docx_path`` to ODT with headless LibreOffice.

    A throw-away user profile is used so a LibreOffice the user has open does
    not block the conversion. Returns the ODT path, or ``None`` if LibreOffice
    is not installed.
    """
    soffice = _find_soffice()
    if soffice is None:
        print("LibreOffice not found - ODT skipped")
        return None
    profile = Path(tempfile.gettempdir()) / "visualopenems_lo_profile"
    profile.mkdir(parents=True, exist_ok=True)
    cmd = [
        soffice,
        f"-env:UserInstallation=file:///{profile.as_posix()}",
        "--headless",
        "--convert-to",
        "odt",
        "--outdir",
        str(docx_path.parent),
        str(docx_path),
    ]
    odt = docx_path.with_suffix(".odt")
    # LibreOffice reports success even when it does not overwrite an existing
    # file (locked/read-only), so delete the target first and verify afterwards.
    odt.unlink(missing_ok=True)
    proc = subprocess.run(cmd, capture_output=True, text=True)
    if proc.returncode != 0 or not odt.exists():
        print("ODT conversion failed:", (proc.stderr or proc.stdout).strip()[:400])
        return None
    return odt


def _inline(paragraph, text: str) -> None:
    """Render `code`, **bold**, *italic* into a paragraph."""
    # Split on code, bold, italic
    pattern = re.compile(r"(`[^`]+`|\*\*[^*]+\*\*|\*[^*]+\*)")
    pos = 0
    for m in pattern.finditer(text):
        if m.start() > pos:
            paragraph.add_run(text[pos : m.start()])
        token = m.group(0)
        if token.startswith("`") and token.endswith("`"):
            run = paragraph.add_run(token[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9)
            rpr = run._element.get_or_add_rPr()
            rfonts = rpr.find(qn("w:rFonts"))
            if rfonts is None:
                rfonts = rpr.makeelement(qn("w:rFonts"), {})
                rpr.append(rfonts)
            rfonts.set(qn("w:ascii"), "Consolas")
            rfonts.set(qn("w:hAnsi"), "Consolas")
        elif token.startswith("**") and token.endswith("**"):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        elif token.startswith("*") and token.endswith("*"):
            run = paragraph.add_run(token[1:-1])
            run.italic = True
        else:
            paragraph.add_run(token)
        pos = m.end()
    if pos < len(text):
        paragraph.add_run(text[pos:])


def _add_table(doc: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    n_cols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=n_cols)
    table.style = "Table Grid"
    for i, row in enumerate(rows):
        for j in range(n_cols):
            cell = table.rows[i].cells[j]
            cell.text = ""
            p = cell.paragraphs[0]
            val = row[j] if j < len(row) else ""
            # strip simple markdown in cells
            val = re.sub(r"\*\*(.+?)\*\*", r"\1", val)
            val = re.sub(r"`(.+?)`", r"\1", val)
            run = p.add_run(val)
            if i == 0:
                run.bold = True
    doc.add_paragraph()


def build(md_path: Path | None = None, docx_path: Path | None = None) -> Path:
    """Render one guide: Markdown -> DOCX. Defaults to the German guide."""
    md_path = Path(md_path) if md_path else MD_PATH
    docx_path = Path(docx_path) if docx_path else md_path.with_suffix(".docx")
    md = md_path.read_text(encoding="utf-8")
    lines = md.splitlines()

    doc = Document()
    # Narrow margins slightly for tables
    for section in doc.sections:
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)

    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    i = 0
    n = len(lines)
    while i < n:
        raw = lines[i]
        line = raw.rstrip()
        stripped = line.strip()

        if not stripped:
            i += 1
            continue

        # Horizontal rule
        if stripped in {"---", "***", "___"}:
            doc.add_paragraph("─" * 60)
            i += 1
            continue

        # Code fence
        if stripped.startswith("```"):
            i += 1
            block: list[str] = []
            while i < n and not lines[i].strip().startswith("```"):
                block.append(lines[i])
                i += 1
            i += 1  # skip closing fence
            p = doc.add_paragraph()
            run = p.add_run("\n".join(block) if block else "")
            run.font.name = "Consolas"
            run.font.size = Pt(9)
            run._element.get_or_add_rPr()
            p.paragraph_format.left_indent = Inches(0.25)
            continue

        # Image: ![alt](relative/path) -> centred picture, caption line follows
        m_img = re.match(r"^!\[[^\]]*\]\(([^)]+)\)$", stripped)
        if m_img:
            img_path = (md_path.parent / m_img.group(1)).resolve()
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            if img_path.exists():
                p.add_run().add_picture(str(img_path), width=Inches(6.2))
            else:
                p.add_run(f"[missing image: {m_img.group(1)}]")
            i += 1
            continue

        # Headings
        m = re.match(r"^(#{1,6})\s+(.*)$", stripped)
        if m:
            level = len(m.group(1))
            text = re.sub(r"\*\*(.+?)\*\*", r"\1", m.group(2))
            text = re.sub(r"`(.+?)`", r"\1", text)
            text = re.sub(r"\*(.+?)\*", r"\1", text)
            doc.add_heading(text, level=min(level, 4))
            i += 1
            continue

        # Table
        if stripped.startswith("|") and i + 1 < n and re.match(r"^\|[\s\-:|]+\|$", lines[i + 1].strip()):
            i += 1  # skip separator
            rows: list[list[str]] = []
            while i < n and lines[i].strip().startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                rows.append(cells)
                i += 1
            _add_table(doc, rows)
            continue

        # Checklist / bullet
        if stripped.startswith("- [ ]") or stripped.startswith("- [x]"):
            checked = stripped.startswith("- [x]")
            text = stripped[5:].strip()
            p = doc.add_paragraph(style="List Bullet")
            run = p.add_run(("[x] " if checked else "[ ] ") + re.sub(r"\*\*(.+?)\*\*", r"\1", text))
            if checked:
                run.bold = True
            i += 1
            continue

        if stripped.startswith("- ") or stripped.startswith("* "):
            text = stripped[2:].strip()
            p = doc.add_paragraph(style="List Bullet")
            _inline(p, text)
            i += 1
            continue

        # Ordered list
        m_ord = re.match(r"^(\d+)\.\s+(.*)$", stripped)
        if m_ord:
            p = doc.add_paragraph(style="List Number")
            _inline(p, m_ord.group(2))
            i += 1
            continue

        # Normal paragraph
        p = doc.add_paragraph()
        _inline(p, stripped)
        i += 1

    # Footer note (language follows the source file name)
    rel = md_path.relative_to(ROOT).as_posix() if md_path.is_relative_to(ROOT) else md_path.name
    if md_path.name.endswith("_DE.md"):
        footer = f"Generiert aus {rel} — VisualopenEMS"
    else:
        footer = f"Generated from {rel} — VisualopenEMS"
    doc.add_paragraph()
    p = doc.add_paragraph()
    run = p.add_run(footer)
    run.italic = True
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)

    doc.save(docx_path)
    return docx_path


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description="Render the guides to DOCX (and ODT).")
    ap.add_argument(
        "--md", type=Path, help="Markdown source (default: docs/ANLEITUNG_DE.md)"
    )
    ap.add_argument(
        "--docx", type=Path, help="DOCX output (default: next to the Markdown)"
    )
    ap.add_argument(
        "--odt", action="store_true", help="also convert to ODT with headless LibreOffice"
    )
    ap.add_argument(
        "--all", action="store_true", help="build both the German and the English guide"
    )
    args = ap.parse_args(argv)

    jobs: list[tuple[Path, Path]] = []
    if args.all:
        for lang in ("DE", "EN"):
            md = ROOT / "docs" / f"ANLEITUNG_{lang}.md"
            jobs.append((md, md.with_suffix(".docx")))
    else:
        md = args.md or MD_PATH
        jobs.append((md, args.docx or md.with_suffix(".docx")))

    rc = 0
    for md, docx in jobs:
        if not md.exists():
            print(f"missing source: {md}")
            rc = 1
            continue
        out = build(md, docx)
        print(f"Wrote {out}")
        print(f"Size: {out.stat().st_size} bytes")
        if args.odt:
            odt = convert_to_odt(out)
            if odt is not None:
                print(f"Wrote {odt}")
                print(f"Size: {odt.stat().st_size} bytes")
    return rc


if __name__ == "__main__":
    raise SystemExit(main())
